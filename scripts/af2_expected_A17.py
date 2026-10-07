"""Expected frames for the sheets of Additional file 2 (A17 §4.4).

Each builder rebuilds one sheet from its source files without importing the scripts that
wrote the workbook. The sheet spec (manuscript/checks/table_specs/AF2_<sheet>.yaml) names the
builder and its parameters. Values are compared cell by cell by 69_check_numbers_A17.py.
"""
from __future__ import annotations

import math
import re
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

import numcheck_A17 as NC

ROOT = NC.ROOT
BOOK = ROOT / "manuscript" / "bmc" / "Additional_file_2.xlsx"
FS = "manuscript/bmc/figure_source"
ID_COLUMNS = re.compile(r"sample|patient|specimen", re.I)

METHOD_LABEL = {
    "BASE-Z": "Baseline", "SA-Z": "HostMix-TOO", "SA-pool22": "22-tissue model", "SA-G": "Gated model",
    "PURE-Zs": "Standardized, pure (C = 0.03)", "PURE-Zs-w": "Standardized, pure (C = 0.15)",
    "MIX-Z0": "Unstandardized, mixtures (C = 0.1)",
}
COHORT_LABEL = {
    "blca_iatlas_imvigor210_2017": "IMvigor210", "brca_iatlas_anders_2022": "Anders", "mel_dfci_2019": "DFCI",
    "paad_iatlas_prince_2022": "PRINCE", "GSE50760": "GSE50760", "prad_su2c_2019": "SU2C/PCF",
    "GSE209998": "AURORA US", "GSE41258": "GSE41258", "GSE14018": "GSE14018", "GSE71729": "GSE71729",
    "GSE74685": "GSE74685", "prad_fhcrc": "FHCRC", "all": "Microarray layer", "MET500": "MET500", "POG570": "POG570",
}


def read(path: str) -> pd.DataFrame:
    full = ROOT / path
    if path.endswith(".json"):
        return pd.json_normalize(pd.read_json(full, typ="series").to_dict())
    return pd.read_csv(full, sep="\t")


def translate(value):
    if isinstance(value, str):
        return NC.TRANSLATE.get(value, value)
    return value


def drop_ids(frame: pd.DataFrame) -> pd.DataFrame:
    return frame.drop(columns=[c for c in frame.columns if ID_COLUMNS.search(str(c))])


def display(value, places: int) -> str:
    return f"{Decimal(str(value)).quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP):f}"


def copy(source: str, columns: list | None = None, rename: dict | None = None, query: str | None = None,
         sort: dict | None = None, maps: dict | None = None, add_map: dict | None = None,
         keep_ids: bool = False) -> pd.DataFrame:
    frame = read(source) if keep_ids else drop_ids(read(source))
    for new, (col, mapping) in (add_map or {}).items():
        frame[new] = frame[col].map(METHOD_LABEL if mapping == "METHOD_LABEL" else mapping)
    if query:
        frame = frame.query(query, engine="python")
    if rename:
        frame = frame.rename(columns=rename)
    for col, mapping in (maps or {}).items():
        mapped = frame[col].map(mapping)
        if mapped.isna().any():
            raise NC.SourceError(f"{source}: unmapped values in {col}: {sorted(set(frame[col][mapped.isna()]))}")
        frame[col] = mapped
    if sort:
        frame = frame.sort_values(sort["by"], ascending=sort.get("ascending", True), kind="stable")
    if columns:
        frame = frame[columns]
    return frame.map(translate).reset_index(drop=True)


def sig10(value):
    return None if value is None or (isinstance(value, float) and math.isnan(value)) else float(f"{value:.10g}")


def power_git(source: str, record: str) -> pd.DataFrame:
    """The power table with the commit that added the preregistration recording it (A18)."""
    import subprocess
    out = subprocess.run(["git", "log", "--diff-filter=A", "--format=%H\t%ci", "--", record], cwd=ROOT,
                         capture_output=True, text=True, check=True).stdout.strip().splitlines()
    if len(out) != 1:
        raise NC.SourceError(f"{record}: {len(out)} adding commits")
    commit, when = out[0].split("\t")
    frame = copy(source)
    frame["recorded_in"] = record
    frame["commit"] = commit
    frame["commit_time"] = when
    return frame


def s3_masking(source: str, layer: str | None = None) -> pd.DataFrame:
    frame = read(source).map(translate)
    if layer:
        recount = read(layer)
        frame = frame.loc[frame["analysis"].ne("microarray")].copy()
        for r in recount.itertuples():
            frame.loc[len(frame) + 2000] = {"note": r.note, "analysis": r.analysis, "cohort": r.cohort,
                                            "method": r.method, "n_truth_not_esophagus": r.n_truth_not_esophagus,
                                            "esophagus_absorption": r.esophagus_fraction}
    frame["esophagus_absorption"] = frame["esophagus_absorption"].map(sig10)
    frame = frame.loc[~(frame["cohort"].eq("MET500") & frame["method"].eq("LD-Z"))]
    frame.loc[len(frame) + 1000] = {"note": "not computed", "analysis": "MET500", "cohort": "MET500",
                                    "method": "LD-Z", "n_truth_not_esophagus": None, "esophagus_absorption": None}
    order = {"MET500": 0, "POG570": 1, "aux_rnaseq": 2, "microarray": 3}
    methods = {"BASE-Z": 0, "SA-Z": 1, "LD-Z": 2, "SCOPE": 3, "CUP-AI-Dx": 4}
    frame = frame.assign(_a=frame["analysis"].map(order), _m=frame["method"].map(methods))
    frame = frame.sort_values(["_a", "_m"], kind="stable").drop(columns=["_a", "_m"])
    return frame.reset_index(drop=True)


def s6_meta(meta: str, drawn: str) -> pd.DataFrame:
    raw = read(meta)
    shown = read(drawn)
    layer_name = {"RNA-seq": "RNA-seq", "마이크로어레이": "Microarray"}
    rows = []
    for _, r in raw.iterrows():
        layer = layer_name[r["layer"]]
        if r["status"] == "cohort":
            name = COHORT_LABEL[r["cohort"]]
            hit = shown.loc[shown["layer"].eq(layer) & shown["row"].eq("cohort") & shown["cohort"].eq(name)]
            rows.append({"layer": layer, "row": "cohort", "cohort": name, "cohort_id": r["cohort"],
                         "n_at_risk": r["n"], "b": r["b"], "c": r["c"], "difference": r["d"], "variance": r["v"],
                         "ci_low": hit["ci_low"].iloc[0], "ci_high": hit["ci_high"].iloc[0],
                         "interval": "Wald, recorded variance", "k": None, "I2": None, "I2_percent": None,
                         "tau2": None})
        else:
            rows.append({"layer": layer, "row": "DerSimonian-Laird", "cohort": None, "cohort_id": None,
                         "n_at_risk": None, "b": None, "c": None, "difference": r["estimate"], "variance": None,
                         "ci_low": r["ci_low"], "ci_high": r["ci_high"],
                         "interval": "DerSimonian-Laird random effects", "k": r["k"], "I2": r["I2"],
                         "I2_percent": float(r["I2"]) * 100, "tau2": r["tau2"]})
    frame = pd.DataFrame(rows)
    frame["_l"] = frame["layer"].map({"RNA-seq": 0, "Microarray": 1})
    frame["_r"] = frame["row"].eq("DerSimonian-Laird").astype(int)
    return frame.sort_values(["_l", "_r"], kind="stable").drop(columns=["_l", "_r"]).reset_index(drop=True)


def s8_display(source: str) -> pd.DataFrame:
    frame = read(source)
    cohort = {"MET500": "MET500", "POG570": "POG570", "aux_rnaseq": "Auxiliary RNA-seq"}
    order = ["BASE-Z", "PURE-Zs", "PURE-Zs-w", "MIX-Z0", "SA-Z"]
    rows = []
    for c in cohort:
        for m in order:
            hit = frame.loc[frame["cohort"].eq(c) & frame["method"].eq(m)]
            if len(hit) != 1:
                raise NC.SourceError(f"{source}: {c} {m} gives {len(hit)} rows")
            r = hit.iloc[0]
            rows.append({"cohort": cohort[c], "model": METHOD_LABEL[m], "n": r["n"], "top1": display(r["top1"], 3),
                         "n_at_risk": r["n_at_risk"], "host_rate": display(r["host_rate"], 3),
                         "n_native_truth": r["n_native_truth"],
                         "native_truth_top1": display(r["native_truth_top1"], 3), "status": "post hoc"})
    return pd.DataFrame(rows)


def s9_microarray(metrics: str, restriction: str, shares: str, category: str) -> pd.DataFrame:
    share_model = {"mix": "MIX-Z0", "std": "PURE-Zs", "stdw": "PURE-Zs-w"}
    gap = {"accuracy": "gap, baseline minus HostMix-TOO top-1 in the evaluation set",
           "random_removal": "gap, HostMix-TOO minus baseline mean loss under random removal"}
    rows = []
    for r in read(metrics).itertuples():
        rows.append({"part": "microarray", "contrast": None, "model": r.model, "display": METHOD_LABEL[r.model],
                     "set": r.set, "cohort": COHORT_LABEL[r.cohort], "platform": None, "metric": r.metric,
                     "value": r.value, "n": r.n})
    for r in read(restriction).itertuples():
        for metric in ("restriction_drop", "random_mean", "random_min", "random_max"):
            rows.append({"part": "missing genes", "contrast": None, "model": r.model,
                         "display": METHOD_LABEL[r.model], "set": "evaluation", "cohort": COHORT_LABEL[r.cohort],
                         "platform": r.platform, "metric": metric, "value": getattr(r, metric),
                         "n": 10 if metric.startswith("random") else None})
    for r in read(shares).itertuples():
        model = share_model[r.name.split("_", 1)[1]]
        base = {"part": "share", "contrast": r.contrast, "model": model, "display": METHOD_LABEL[model],
                "set": None, "cohort": None, "platform": None, "n": None}
        rows.append({**base, "metric": "share", "value": r.value})
        rows.append({**base, "metric": gap[r.contrast], "value": r.gap})
    for r in read(category).itertuples():
        rows.append({"part": "reading", "contrast": r.contrast, "model": None, "display": None, "set": None,
                     "cohort": None, "platform": None, "metric": "category", "value": r.category, "n": None})
    return pd.DataFrame(rows)


def fig4a(source: str, files: dict) -> pd.DataFrame:
    frame = read(source).rename(columns={"group": "role", "group_display": "group"})
    frame = frame[["group", "role", "cohort", "method", "host_rate", "n_at_risk", "source"]].copy()
    frame["source_file"] = frame["source"].map(files)
    column = {"results/stage8/external/tables/metrics.tsv": "analysis_cohort"}
    frame["status"] = [NC.status_of(f, f"{column.get(f, 'cohort')} == '{c}' and method == '{m}'")
                       for f, c, m in zip(frame["source_file"], frame["cohort"], frame["method"])]
    return frame[["group", "role", "cohort", "method", "host_rate", "n_at_risk", "source", "status", "source_file"]]


BUILDERS = {"copy": copy, "s3_masking": s3_masking, "s6_meta": s6_meta, "s8_display": s8_display,
            "s9_microarray": s9_microarray, "fig4a": fig4a, "power_git": power_git}


def sheet_frame(name: str, header_row: int) -> tuple[str, pd.DataFrame]:
    book = load_workbook(BOOK, read_only=True)
    rows = list(book[name].iter_rows(values_only=True))
    title = rows[0][0]
    header = [str(c) for c in rows[header_row - 1]]
    while header and header[-1] == "None":
        header.pop()
    body = [list(r[:len(header)]) for r in rows[header_row:]]
    return title, pd.DataFrame(body, columns=header)


def same(a, b) -> bool:
    def empty(x):
        return x is None or (isinstance(x, float) and math.isnan(x)) or (x is pd.NA) or x == ""
    if empty(a) and empty(b):
        return True
    if empty(a) or empty(b):
        return False
    if isinstance(a, bool) or isinstance(b, bool):
        return str(a) == str(b)
    if isinstance(a, (int, float)) and not isinstance(b, str):
        try:
            fb = float(b)
        except (TypeError, ValueError):
            return False
        return abs(float(a) - fb) <= 1e-12 * max(1.0, abs(fb))
    return str(a) == str(b)
