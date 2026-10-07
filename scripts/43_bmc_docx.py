"""Write BMC DOCX files with python-docx.

pandoc is not installed on this machine, and the package index could not
be reached. Line spacing, line numbers, and page numbers are set here.
"""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]


def add_line_numbers(section) -> None:
    ln = OxmlElement("w:lnNumType")
    ln.set(qn("w:countBy"), "1")
    ln.set(qn("w:restart"), "continuous")
    section._sectPr.append(ln)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def style_paragraph(paragraph) -> None:
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 2.0
    fmt.space_after = Pt(0)
    for run in paragraph.runs:
        run.font.name = "Liberation Serif"
        run.font.size = Pt(12)


def add_table(document: Document, lines: list[str]) -> None:
    rows = []
    for line in lines:
        if re.match(r"^\|?\s*:?-{3,}", line.replace("|", " ").strip()) or set(line.replace("|", "").replace(":", "").replace("-", "").strip()) == set():
            continue
        if re.match(r"^\s*\|?\s*-+", line):
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        if cells and not all(set(cell) <= set("-: ") for cell in cells):
            rows.append(cells)
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(width):
            table.rows[i].cells[j].text = row[j] if j < len(row) else ""


def add_runs(paragraph, text: str) -> None:
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)


def add_markdown(document: Document, text: str) -> None:
    blocks = re.split(r"\n\s*\n", text.strip())
    for block in blocks:
        lines = block.strip("\n")
        if not lines.strip():
            continue
        content = [line for line in lines.splitlines() if line.strip()]
        if content and all(line.strip().startswith("|") for line in content):
            add_table(document, content)
            continue
        if lines.startswith("# "):
            paragraph = document.add_heading(lines[2:].replace("**", "").strip(), level=1)
        elif lines.startswith("## "):
            paragraph = document.add_heading(lines[3:].replace("**", "").strip(), level=2)
        elif lines.startswith("### "):
            paragraph = document.add_heading(lines[4:].replace("**", "").strip(), level=3)
        else:
            paragraph = document.add_paragraph()
            add_runs(paragraph, lines.replace("\n", " "))
        style_paragraph(paragraph)


def write_docx(src: Path, dest: Path, images: list[Path] | None = None) -> None:
    document = Document()
    section = document.sections[0]
    add_line_numbers(section)
    add_page_number(section.footer.paragraphs[0])
    add_markdown(document, src.read_text())
    for image in images or []:
        if image.exists():
            document.add_picture(str(image), width=Inches(6.5))
    dest.parent.mkdir(parents=True, exist_ok=True)
    document.save(dest)
    print(dest, dest.stat().st_size)


def main() -> None:
    figures = ROOT / "manuscript" / "bmc" / "figures"
    write_docx(ROOT / "manuscript" / "bmc" / "A_main.md", ROOT / "manuscript" / "bmc" / "A_main.docx")
    write_docx(
        ROOT / "manuscript" / "bmc" / "Additional_file_1.md",
        ROOT / "manuscript" / "bmc" / "Additional_file_1.docx",
        [figures / f"FigureS{index}.png" for index in range(1, 7)],
    )


if __name__ == "__main__":
    main()
