"""A17 Part A: status of three Additional file 2 sheets and the reference evidence table.

Fig6a_cohorts gains a per-row status column. The status follows the source rule of A14
section 2.3 with results/stage10/ added as post hoc: results/stage8/external/ is
pre-specified descriptive, except the POG570 at-risk rates, which are the two arms of H1
in the confirmatory set; results/stage9/ and new aggregates of stored predictions are
post hoc; results/stage7/confirm/tables/metrics.tsv is pre-specified descriptive.
Sensitivity and Conformal are named in the preregistrations, so both are pre-specified
descriptive. Every other sheet is left as it is.
"""
from __future__ import annotations

import importlib
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
B58 = importlib.import_module("58_additional2_figures_A15")

BOOK = B58.BOOK
FIG6A = ("Cohort host-attraction rates drawn in Figure 6a, with the group labels of the figure. "
         "Status is given per row in the column 'status'.")
SENSITIVITY = ("Pre-specified descriptive auxiliary sensitivity rows (auxiliary preregistration, descriptive "
               "analysis 5). Effect sizes only.")
CONFORMAL = ("Pre-specified descriptive POG570 conformal summary (stage-3 preregistration, supporting "
             "analyses). Aggregate only.")
STATUS = {
    ("stage8 metrics", "Development"): "pre-specified descriptive",
    ("stage8 metrics", "Confirmation 1"): "preregistered",
    ("subcohort_metrics", "Confirmation 2"): "post hoc",
    ("stage7 metrics", "Descriptive RNA-seq"): "pre-specified descriptive",
    ("post hoc aggregate of confirmatory predictions", "Descriptive microarray"): "post hoc",
}
SOURCE_FILE = {
    "stage8 metrics": "results/stage8/external/tables/metrics.tsv",
    "subcohort_metrics": "results/stage9/subcohort_metrics.tsv",
    "stage7 metrics": "results/stage7/confirm/tables/metrics.tsv",
    "post hoc aggregate of confirmatory predictions": "results/stage7/confirm/predictions.parquet",
}

PUB = "data/reference/publisher_A17"
REFS = [
    {"number": 23, "field": "pages", "manuscript": "pl1", "crossref": "no page and no article number",
     "evidence": "Science Signaling 2013;6(269):pl1 (PubMed 23550210: volume 6, issue 269, pages pl1)",
     "publisher_page": "https://www.science.org/doi/10.1126/scisignal.2004088",
     "publisher_result": "not opened: HTTP 403 on 2026-10-07 00:24 UTC",
     "other_source": f"PubMed esummary 23550210, {PUB}/pubmed_23_26.json",
     "decision": "kept"},
    {"number": 26, "field": "year", "manuscript": "2013;41(D1)", "crossref": "issued 2012-11-26 (online only)",
     "evidence": "Database issue of January 2013 (PubMed 23193258: pubdate 2013 Jan, epubdate 2012 Nov 27, "
                 "volume 41, issue 'Database issue', pages D991-5)",
     "publisher_page": "https://academic.oup.com/nar/article/41/D1/D991/1067995",
     "publisher_result": "not opened: HTTP 403 on 2026-10-07 00:24 UTC",
     "other_source": f"PubMed esummary 23193258, {PUB}/pubmed_23_26.json",
     "decision": "kept"},
    {"number": 40, "field": "year, volume, issue, pages", "manuscript": "2023;4(1):128–147",
     "crossref": "issued 2022-12-30 (online), no volume, issue or pages",
     "evidence": "citation tags of the article page: citation_volume 4, citation_issue 1, citation_firstpage 128, "
                 "citation_lastpage 147, citation_publication_date 2023/01, citation_online_date 2022/12/30",
     "publisher_page": "https://www.nature.com/articles/s43018-022-00491-x",
     "publisher_result": "opened 2026-10-07 00:24 UTC (HTTP 200)",
     "other_source": f"{PUB}/ref40.html",
     "decision": "kept"},
    {"number": 54, "field": "journal abbreviation", "manuscript": "J Mach Learn Res", "crossref": "no DOI",
     "evidence": "journal site: electronic ISSN 1533-7928, paper volumes ISSN 1532-4435; J_Medline.txt JrId 32661 "
                 "MedAbbr 'J Mach Learn Res', ISSN (Print) 1532-4435, ISSN (Online) 1533-7928, NlmId 101262635; "
                 "NLM Catalog 101262635 medlineta 'J Mach Learn Res'",
     "publisher_page": "https://www.jmlr.org/ and https://www.jmlr.org/papers/v12/pedregosa11a.html "
                       "(citation_issn 1533-7928, volume 12, issue 85, pages 2825-2830, 2011)",
     "publisher_result": "opened 2026-10-07 00:23 UTC (HTTP 200)",
     "other_source": f"data/reference/nlm/J_Medline.txt, {PUB}/nlmcatalog_101262635.json",
     "decision": "J Mach Learn Res, source J_Medline MedAbbr by ISSN"},
]


def main() -> None:
    book = load_workbook(BOOK)
    sheet = book["Fig6a_cohorts"]
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    frame = pd.DataFrame(rows[1:], columns=rows[0])
    if "status" not in frame.columns:
        keys = list(zip(frame["source"], frame["role"]))
        missing = sorted(set(keys) - set(STATUS))
        if missing:
            raise SystemExit(f"Fig6a: no status rule for {missing}")
        frame["status"] = [STATUS[key] for key in keys]
        frame["source_file"] = frame["source"].map(SOURCE_FILE)
    B58.DESCRIPTION["Fig6a_cohorts"] = FIG6A
    B58.DESCRIPTION["Sensitivity"] = SENSITIVITY
    B58.DESCRIPTION["Conformal"] = CONFORMAL
    B58.put(book, "Fig6a_cohorts", frame)
    book["Sensitivity"].cell(1, 1).value = SENSITIVITY
    book["Conformal"].cell(1, 1).value = CONFORMAL
    B58.readme(book)
    book.save(BOOK)
    print(frame.groupby(["role", "status"]).size().to_string())
    out = ROOT / "manuscript" / "checks" / "refs_sources_A17.tsv"
    pd.DataFrame(REFS).to_csv(out, sep="\t", index=False)
    print(out, len(REFS))


if __name__ == "__main__":
    main()
