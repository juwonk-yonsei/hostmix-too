#!/usr/bin/env python3
"""Document properties (creator, last modified by, title) of Additional files 2 and 3.

The workbooks are edited in place by several earlier scripts, so only the properties are set here; cell values
and sheet order are not touched. The DOCX files get their properties in 80_submission_A18.py.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
BMC = ROOT / "manuscript" / "bmc"
AUTHORS = "Juwon Kang; Junjeong Choi"
TITLES = {
    "Additional_file_2.xlsx": "Additional file 2. Source values for the figures and tables",
    "Additional_file_3.xlsx": "Additional file 3. Internal preregistration history and deviations",
}


def main() -> None:
    for name, title in TITLES.items():
        path = BMC / name
        book = load_workbook(path)
        before = [[c.value for c in row] for ws in book.worksheets for row in ws.iter_rows()]
        props = book.properties
        print(name, "before:", props.creator, props.lastModifiedBy, props.title)
        props.creator = AUTHORS
        props.lastModifiedBy = AUTHORS
        props.title = title
        book.save(path)
        after_book = load_workbook(path)
        after = [[c.value for c in row] for ws in after_book.worksheets for row in ws.iter_rows()]
        if after != before or after_book.sheetnames != book.sheetnames:
            raise SystemExit(f"{name}: cell values or sheets changed")
        p = after_book.properties
        print(name, "after: ", p.creator, p.lastModifiedBy, p.title)


if __name__ == "__main__":
    main()
