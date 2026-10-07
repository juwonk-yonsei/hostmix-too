#!/usr/bin/env python3
"""A19 section 3.5: layer-set note in row 1 of the Additional file 2 sheet S3_esophagus_masking.

The sentence is appended to row 1 and to the README text of the sheet, which must stay equal. The sheet S8_shares
is put back directly after Fig5_ablation, where it stood next to the ablation sheet before the renumbering;
the README rows follow the sheet order.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT / "manuscript" / "bmc" / "Additional_file_2.xlsx"
SHEET = "S3_esophagus_masking"
NOTE = "Microarray rows use the layer set (one sample per patient across the layer, n = 536)."


def main() -> None:
    book = load_workbook(PATH)
    ws = book[SHEET]
    title = ws.cell(1, 1).value
    if NOTE in title:
        raise SystemExit("note already present")
    new_title = f"{title.rstrip()} {NOTE}"
    ws.cell(1, 1).value = new_title
    readme = book["README"]
    hits = [row for row in readme.iter_rows(min_row=2) if row[0].value == SHEET]
    if len(hits) != 1 or hits[0][1].value != title:
        raise SystemExit("README row of the sheet not found or not equal to row 1")
    hits[0][1].value = new_title

    sheets = [s for s in book.worksheets if s.title != "S8_shares"]
    at = [s.title for s in sheets].index("Fig5_ablation") + 1
    sheets.insert(at, book["S8_shares"])
    book._sheets = sheets
    rows = [[c.value for c in r] for r in readme.iter_rows(min_row=2)]
    position = {s.title: i for i, s in enumerate(sheets)}
    rows.sort(key=lambda r: position.get(r[0], len(position)))
    for i, values in enumerate(rows, start=2):
        for j, value in enumerate(values, start=1):
            readme.cell(i, j).value = value
    book.save(PATH)
    print("row 1 before:", title)
    print("row 1 after: ", new_title)
    print("sheets:", [s.title for s in book.worksheets])


if __name__ == "__main__":
    main()
